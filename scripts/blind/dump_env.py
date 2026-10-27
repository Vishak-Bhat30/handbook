"""Render a task's initial environment (never its tests/) as text. Usage: dump_env.py <task_dir> <out.txt>"""
import json
import subprocess
import sys
from pathlib import Path

import openpyxl

task = Path(sys.argv[1])
env = task / "environment"
out = []
out.append("#### INSTRUCTION\n" + (task / "instruction.md").read_text())
out.append("#### SYSTEM PROMPT\n" + (task / "system_prompt.md").read_text())
svc = env / "initial_external_services"
mail = svc / "google_mail" / "inbox.json"
if mail.exists():
    m = json.loads(mail.read_text())
    out.append(f"#### MAILBOX {m.get('mailbox')} folders={[f['name'] for f in m.get('folders', [])]}")
    for e in sorted(m["emails"], key=lambda e: e["date"]):
        out.append(f"--- [{e['email_id']}] {e['folder']} {e['date']} read={e.get('is_read')}\nFrom: {e['from_addr']}\nTo: {e['to_addr']}\nCc: {e.get('cc_addr')}\n"
                   f"Subject: {e['subject']}\nAttachments: {e.get('attachments')}\n{e.get('body_text')}")
sl = svc / "slack" / "slack.json"
if sl.exists():
    s = json.loads(sl.read_text())
    users = {u: d.get("real_name") for u, d in s["users"].items()}
    out.append("#### SLACK USERS\n" + "\n".join(f"{u} {d.get('name')} {d.get('real_name')} ({d['profile'].get('title')})" for u, d in s["users"].items()))
    chans = s["channels"]
    out.append("#### SLACK CHANNELS\n" + json.dumps({c: (d.get("name"), d.get("is_im"), d.get("members")) for c, d in chans.items()}))
    msgs = s["messages"]
    for c, ms in msgs.items():
        out.append(f"#### SLACK #{chans.get(c, {}).get('name', c)} ({c})")
        for mm in sorted(ms, key=lambda x: float(x.get("ts", 0))):
            import datetime
            t = datetime.datetime.utcfromtimestamp(float(mm["ts"])).isoformat()
            out.append(f"[{t}Z ts={mm['ts']} thread={mm.get('thread_ts')}] {mm.get('user')}({users.get(mm.get('user'))}): {mm.get('text')}")
cal = svc / "google_calendar" / "calendar_data.json"
if cal.exists():
    c = json.loads(cal.read_text())
    out.append("#### CALENDAR")
    evs = c.get("events", [])
    for ev in (evs.values() if isinstance(evs, dict) else evs):
        if isinstance(ev, str):
            out.append(ev); continue
        out.append(json.dumps({k: ev.get(k) for k in ("id", "summary", "start", "end", "description", "attendees", "location")}))
jr = svc / "jira" / "jira_state.json"
if jr.exists():
    j = json.loads(jr.read_text())

    def adf(node):
        if isinstance(node, str):
            return node
        if isinstance(node, dict):
            return (node.get("text") or "") + "".join(adf(c) for c in node.get("content", [])) + (
                "\n" if node.get("type") == "paragraph" else "")
        if isinstance(node, list):
            return "".join(adf(c) for c in node)
        return ""
    out.append("#### JIRA users " + json.dumps({k: (u.get("displayName"), u.get("emailAddress")) for k, u in j.get("users", {}).items()}))
    out.append("#### JIRA statuses " + json.dumps([s["name"] for s in j.get("statuses", {}).values()]))
    out.append("#### JIRA fields " + json.dumps([(f.get("id"), f.get("name")) for f in j.get("fields", [])]))
    for k, iss in j.get("issues", {}).items():
        f = iss.get("fields", {})
        meta = {x: (f.get(x).get("name") or f.get(x).get("displayName") or f.get(x).get("key")) if isinstance(f.get(x), dict) else f.get(x)
                for x in f if x not in ("description", "summary")}
        out.append(f"#### JIRA {k}: {f.get('summary')}\n{json.dumps(meta, default=str)}\n{adf(f.get('description'))}")
    for k, cs in (j.get("comments") or {}).items():
        for c in cs:
            out.append(f"#### JIRA COMMENT {k} by {(c.get('author') or {}).get('displayName')}: {adf(c.get('body'))}")
    if j.get("linkTypes"):
        out.append("#### JIRA linkTypes " + json.dumps([l.get("name") for l in j["linkTypes"]]))
import tempfile
import zipfile

ws_files = sorted((env / "initial_workspace").iterdir())
for z in [f for f in ws_files if f.suffix == ".zip"]:
    tmp = Path(tempfile.mkdtemp())
    zipfile.ZipFile(z).extractall(tmp)
    ws_files += sorted(p for p in tmp.rglob("*") if p.is_file() and not p.name.startswith("."))
    out.append(f"#### ZIP {z.name}: {[str(p.relative_to(tmp)) for p in tmp.rglob('*') if p.is_file()]}")
for f in ws_files:
    if f.suffix == ".zip" or "__MACOSX" in str(f):
        continue
    if f.suffix == ".xlsx":
        wb = openpyxl.load_workbook(f, data_only=False)
        for ws in wb.worksheets:
            out.append(f"#### XLSX {f.name} :: {ws.title} ({ws.max_row}x{ws.max_column})")
            for row in ws.iter_rows(values_only=True):
                if any(v is not None for v in row):
                    out.append(" | ".join("" if v is None else str(v) for v in row))
    elif f.suffix == ".pdf" and "SOP" not in f.name:
        import pdfplumber
        with pdfplumber.open(f) as pdf:
            txt = "\n".join(pg.extract_text() or "" for pg in pdf.pages)
        out.append(f"#### PDF {f.name}\n{txt}")
    elif "SOP" not in f.name:
        out.append(f"#### FILE {f.name}")
Path(sys.argv[2]).write_text("\n".join(out))
