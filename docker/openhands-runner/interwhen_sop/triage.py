"""Step 2 of scripts/algo.md: per-paragraph triage of a task's SOP.

    python -m interwhen_sop.triage init <task>                 # deterministic pre-pass (headings, cited paragraphs)
    python -m interwhen_sop.triage pending <task>              # paragraphs still waiting for a builder verdict
    python -m interwhen_sop.triage apply <task> <file.tsv>     # builder verdicts: id<TAB>verdict<TAB>verifiers<TAB>reason
    python -m interwhen_sop.triage propose <task> [--workers]  # optional: SLM drafts verdicts (slow, needs review)
    python -m interwhen_sop.triage check <task>                # coverage gate (exit 1 on gaps)
    python -m interwhen_sop.triage report <task>               # verdict counts

Writes interwhen_sop/tasks/<task>/triage.json: one entry per SOP paragraph with
  verdict  context | data | not_observable | not_applicable | verifiable
  trigger, stage (pre|completion|both), oracle, strictness (exact|semantic), check, reason
  verifiers  rule ids implementing it, reviewed (builder confirmed the verdict)
The builder decides every non-trivial verdict (algo.md Step 2); SLM drafts are only suggestions.
Inputs are the SOP, instruction, system prompt and seed data - never tests/.
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path

from . import framework, slm, sopdata

REPO = Path(__file__).resolve().parents[3]
VERDICTS = ("context", "data", "not_observable", "not_applicable", "verifiable")
TOOLS = """Tools the agent has (every call is visible to the verifier, with arguments and output):
- files in /workdir: list/read files, read PDF/DOCX, run Python/Bash (create/edit .xlsx/.pdf/.docx, copy, unzip)
- google_mail: search/read emails, send_email(to, cc, subject, body, attachments), reply, forward_email(email_id, to, body)
- slack: list channels/users, read history, post_message(channel_id, text), DM
- google_calendar: list/create/update events (summary, start, end, description, attendees)
- jira: search/create/update/transition issues, comment
Not available: phone calls, fax, physical actions, waiting for future days, other staff's actions."""

PROMPT = """You are helping build runtime verifiers for an AI office agent that must follow a company SOP.
Decide whether compliance with ONE SOP paragraph can be checked from the agent's tool calls (arguments, outputs)
or from the end state of its files/mail/Slack/calendar, in THIS task.

Verdicts:
- context: heading, purpose, definition, governance; no instruction.
- data: a lookup other rules use (template text, code table row, contact, column list).
- not_observable: phone calls, physical handling, real-time waiting, or other people's actions.
- not_applicable: an instruction whose trigger cannot occur given the task facts below (e.g. a rule for a
  denial code that appears nowhere in the facts). General prohibitions are always applicable.
- verifiable: an instruction whose compliance shows up in tool calls or end state in this task.

{tools}

TASK FACTS (instruction, date, files, mail, Slack - what the agent can see):
{facts}

SOP section {section} "{title}". Neighbouring paragraphs:
{before}
>>> PARAGRAPH {pid}: {text}
{after}
{refs}
Reply with one JSON object:
{{"verdict": "...", "reason": "<one sentence>", "trigger": "<which tool calls/files/messages it governs or null>",
 "stage": "pre|completion|both|null", "oracle": "<where the expected value comes from or null>",
 "strictness": "exact|semantic|null", "check": "<the concrete check in one sentence or null>"}}"""


def task_dir(task: str) -> Path:
    return REPO / "tasks" / task


def triage_path(task: str) -> Path:
    return Path(__file__).parent / "tasks" / task / "triage.json"


def _xlsx_summary(path: Path) -> str:
    try:
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        out = []
        for ws in wb.worksheets[:2]:
            rows = [[("" if c is None else str(c)) for c in r] for r in ws.iter_rows(values_only=True)]
            rows = [r for r in rows if any(x.strip() for x in r)]
            head = " | ".join(x for x in (rows[0] if rows else []) if x)[:400]
            first = " | ".join(rows[1])[:300] if len(rows) > 1 else ""
            out.append(f"[{ws.title}] {len(rows)} rows; header: {head}" + (f"; row1: {first}" if first else ""))
        return " ".join(out)
    except Exception as exc:  # unreadable workbook: list it anyway
        return f"(unreadable: {type(exc).__name__})"


def _pdf_text(path: Path, n: int = 700) -> str:
    try:
        import pdfplumber
        with pdfplumber.open(path) as pdf:
            txt = " ".join(pg.extract_text() or "" for pg in pdf.pages[:3])
        return re.sub(r"\s+", " ", txt)[:n]
    except Exception as exc:
        return f"(unreadable: {type(exc).__name__})"


def _zip_list(path: Path) -> str:
    import zipfile
    try:
        names = [n for n in zipfile.ZipFile(path).namelist() if "__MACOSX" not in n and not n.endswith("/")]
        return "contains " + ", ".join(names[:25])
    except Exception as exc:
        return f"(unreadable: {type(exc).__name__})"


def facts(task: str, limit: int = 22000) -> str:
    d = task_dir(task)
    env = d / "environment"
    out = ["INSTRUCTION: " + (d / "instruction.md").read_text().strip(),
           "SYSTEM PROMPT: " + (d / "system_prompt.md").read_text().strip()[:600], "FILES:"]
    ws = env / "initial_workspace"
    for p in sorted(ws.rglob("*")):
        if p.is_file() and "__MACOSX" not in p.parts and not p.name.startswith(("._", "~$")):
            rel, suf = p.relative_to(ws), p.suffix.lower()
            info = (_xlsx_summary(p) if suf == ".xlsx" else _zip_list(p) if suf == ".zip"
                    else _pdf_text(p) if suf == ".pdf" and not re.search(r"SOP", p.name, re.I) else "")
            out.append(f"- {rel}" + (f": {info}" if info else ""))
    svc = env / "initial_external_services"
    mail = svc / "google_mail" / "inbox.json"
    if mail.exists():
        m = json.loads(mail.read_text())
        out.append(f"MAILBOX of {m.get('mailbox')}:")
        for e in sorted(m.get("emails", []), key=lambda e: e.get("date", "")):
            atts = [a.get("filename") for a in e.get("attachments") or []]
            body = re.sub(r"\s+", " ", e.get("body_text") or "")[:350]
            out.append(f"- [{e.get('folder')}] {e.get('date')} from {e.get('from_addr')} to {e.get('to_addr')} | "
                       f"{e.get('subject')} | att {atts} | {body}")
    sl = svc / "slack" / "slack.json"
    if sl.exists():
        s = json.loads(sl.read_text())
        chans = s.get("channels", {})
        out.append("SLACK channels: " + ", ".join(f"#{c.get('name')}" for c in chans.values()))
        for c, ms in s.get("messages", {}).items():
            for mm in sorted(ms, key=lambda x: float(x.get("ts", 0)))[-6:]:
                out.append(f"- #{chans.get(c, {}).get('name', c)}: {(mm.get('text') or '')[:200]}")
    cal = svc / "google_calendar" / "calendar_data.json"
    if cal.exists():
        evs = json.loads(cal.read_text()).get("events", [])
        out.append(f"CALENDAR: {len(evs)} events")
    text = "\n".join(out)
    return text if len(text) <= limit else text[:limit] + "\n...(truncated)"


def _context(sop, i: int) -> tuple[str, str, str]:
    ps = sop.paragraphs
    sec = ps[i]["section"]
    near = lambda js: "\n".join(f"  [{ps[j]['id']} {ps[j]['section']}] {ps[j]['text'][:400]}" for j in js
                                if 0 <= j < len(ps) and ps[j]["section"].split(".")[0] == sec.split(".")[0])
    refs = []
    for m in set(re.findall(r"Template\s+(\d+)", ps[i]["text"])):
        hit = next((p for p in ps if re.match(rf"Template\s+{m}\b", p["text"])), None)
        if hit:
            idx = ps.index(hit)
            refs.append("Referenced template:\n" + "\n".join("  " + p["text"][:300] for p in ps[idx:idx + 6]))
    for m in sorted(set(re.findall(r"Section\s+(\d+(?:\.\d+)?)", ps[i]["text"])))[:2]:
        sub = [p for p in ps if p["section"] == m and p["kind"] != "heading"][:3]
        if sub:
            refs.append(f"Referenced section {m}:\n" + "\n".join("  " + p["text"][:300] for p in sub))
    return near(range(i - 2, i)), near(range(i + 1, i + 3)), "\n".join(refs)


def _ask(prompt: str) -> dict | None:
    model = slm.model_name()
    if model is None:
        return None
    try:
        resp = slm._post("/chat/completions", {
            "model": model, "temperature": 0, "max_tokens": 6000,
            "messages": [{"role": "user", "content": prompt}]}, 300)
        return slm.parse_json_answer(resp["choices"][0]["message"].get("content") or "")
    except Exception:
        return None


def _load(task: str) -> dict[str, dict]:
    path = triage_path(task)
    return {e["id"]: e for e in json.loads(path.read_text())["paragraphs"]} if path.exists() else {}


def _save(task: str, entries: dict[str, dict]) -> None:
    sop = sopdata.load(task)
    rows = [entries.get(p["id"]) or {"id": p["id"], "section": p["section"], "kind": p["kind"], "text": p["text"],
                                    "verdict": None} for p in sop.paragraphs]
    triage_path(task).write_text(json.dumps({"task": task, "paragraphs": rows}, indent=1, ensure_ascii=False))


def init(task: str) -> None:
    """Headings/labels and document metadata -> context; paragraphs the verifier module cites -> verifiable.
    Reviewed entries are kept."""
    sop = sopdata.load(task)
    old = _load(task)
    cited = cited_paragraphs(task)
    out = {}
    for p in sop.paragraphs:
        base = {"id": p["id"], "section": p["section"], "kind": p["kind"], "text": p["text"]}
        prev = old.get(p["id"])
        if prev and prev.get("reviewed") and prev.get("text") == p["text"]:
            out[p["id"]] = prev
        elif p["kind"] == "heading" or len(p["text"].split()) < 4 or p["section"] == "0":
            out[p["id"]] = {**base, "verdict": "context", "reason": "heading / label / document metadata",
                            "source": "auto", "reviewed": True}
        elif p["id"] in cited:
            out[p["id"]] = {**base, "verdict": "verifiable", "verifiers": sorted(set(cited[p["id"]])),
                            "reason": "cited by the task's verifier rules", "source": "cited", "reviewed": True}
        else:
            out[p["id"]] = {**base, "verdict": None}
    _save(task, out)


def pending(task: str) -> None:
    for e in _load(task).values():
        if not e.get("reviewed"):
            print(f"{e['id']}\t{e['section']}\t{e['kind']}\t{e['text']}")


def apply(task: str, tsv: str) -> None:
    entries = _load(task)
    n = 0
    for line in Path(tsv).read_text().splitlines():
        if not line.strip() or line.startswith("#"):
            continue
        parts = (line.split("\t") + ["", "", ""])[:4]
        pid, verdict, verifiers, reason = (x.strip() for x in parts)
        if pid not in entries or verdict not in VERDICTS:
            print(f"skip: {line[:80]}", file=sys.stderr)
            continue
        entries[pid].update(verdict=verdict, reason=reason, source="builder", reviewed=True,
                            verifiers=[v for v in verifiers.split(",") if v.strip()])
        n += 1
    _save(task, entries)
    print(f"{task}: applied {n} verdicts")


def propose(task: str, workers: int = 16, redo: bool = False) -> dict:
    sop = sopdata.load(task)
    path = triage_path(task)
    old = {e["id"]: e for e in json.loads(path.read_text())["paragraphs"]} if path.exists() else {}
    fact = facts(task)
    entries: dict[str, dict] = {}
    todo = []
    for i, p in enumerate(sop.paragraphs):
        base = {"id": p["id"], "section": p["section"], "kind": p["kind"], "text": p["text"]}
        prev = old.get(p["id"])
        if prev and prev.get("text") == p["text"] and (prev.get("reviewed") or (not redo and prev.get("verdict"))):
            entries[p["id"]] = prev
        elif p["kind"] == "heading" or len(p["text"].split()) < 4:
            entries[p["id"]] = {**base, "verdict": "context", "reason": "heading / label", "source": "auto"}
        else:
            todo.append((i, base))
    print(f"{task}: {len(sop.paragraphs)} paragraphs, {len(todo)} to triage with the SLM", file=sys.stderr)

    def one(i, base):
        before, after, refs = _context(sop, i)
        p = sop.paragraphs[i]
        ans = _ask(PROMPT.format(tools=TOOLS, facts=fact, section=p["section"], title=p.get("section_title", ""),
                                 before=before, pid=p["id"], text=p["text"], after=after, refs=refs))
        if not ans or ans.get("verdict") not in VERDICTS:
            return {**base, "verdict": None, "reason": "SLM gave no valid verdict", "source": "slm"}
        keep = ("verdict", "reason", "trigger", "stage", "oracle", "strictness", "check")
        return {**base, **{k: (None if ans.get(k) in ("null", "") else ans.get(k)) for k in keep}, "source": "slm"}

    def save():
        rows = [entries.get(p["id"]) or old.get(p["id"]) or {"id": p["id"], "section": p["section"], "kind": p["kind"],
                                                             "text": p["text"], "verdict": None}
                for p in sop.paragraphs]
        path.write_text(json.dumps({"task": task, "paragraphs": rows}, indent=1, ensure_ascii=False))

    with ThreadPoolExecutor(workers) as ex:
        futs = [ex.submit(one, i, b) for i, b in todo]
        for n, f in enumerate(as_completed(futs), 1):
            e = f.result()
            entries[e["id"]] = {**e, "verifiers": (old.get(e["id"]) or {}).get("verifiers", [])}
            if n % 25 == 0:
                print(f"  {n}/{len(todo)}", file=sys.stderr)
                save()
    save()
    return json.loads(path.read_text())


def cited_paragraphs(task: str) -> dict[str, list[str]]:
    """Paragraph ids cited by the task's verifier module (and the shared helpers it uses), via their
    (section, paragraph prefix) reference tuples."""
    sop = sopdata.load(task)
    mod = sys.modules.get(framework.verifier_module(task)) or __import__(framework.verifier_module(task),
                                                                          fromlist=["RULES"])
    srcs = [Path(mod.__file__)] + [Path(__file__).parent / f for f in ("kit.py",)]
    out: dict[str, list[str]] = {}
    for src in srcs:
        for sec, prefix in re.findall(r"\(\s*\"(\d+(?:\.\d+)*)\"\s*,\s*\"([^\"]+)\"", src.read_text()):
            try:
                out.setdefault(sop.find(sec, prefix), []).append(f"module:{src.stem}")
            except KeyError:
                continue
    return out


def check(task: str) -> int:
    data = json.loads(triage_path(task).read_text())
    ids = {r.id for r in framework.load_rules(task)}
    cited = cited_paragraphs(task)
    problems = []
    for e in data["paragraphs"]:
        if e.get("verdict") not in VERDICTS:
            problems.append(f"{e['id']}: no verdict")
        elif e["verdict"] == "verifiable":
            bad = [v for v in e.get("verifiers", []) if not v.startswith("module:") and v.split("#")[0] not in ids]
            if bad:
                problems.append(f"{e['id']}: unknown rule ids {bad}")
            elif not e.get("verifiers") and e["id"] not in cited:
                problems.append(f"{e['id']} [{e['section']}] verifiable but no verifier: {e['text'][:110]}")
    for p in problems:
        print(p)
    n = len(data["paragraphs"])
    print(f"{task}: {n} paragraphs, {len(problems)} gaps")
    return 1 if problems else 0


def report(task: str) -> None:
    data = json.loads(triage_path(task).read_text())
    cited = cited_paragraphs(task)
    counts: dict[str, int] = {}
    for e in data["paragraphs"]:
        counts[e.get("verdict") or "none"] = counts.get(e.get("verdict") or "none", 0) + 1
    ver = [e for e in data["paragraphs"] if e.get("verdict") == "verifiable"]
    done = [e for e in ver if e.get("verifiers") or e["id"] in cited]
    print(json.dumps({"task": task, "verdicts": counts, "verifiable": len(ver), "implemented": len(done)}))


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("cmd", choices=("init", "pending", "apply", "propose", "check", "report"))
    ap.add_argument("task")
    ap.add_argument("file", nargs="?")
    ap.add_argument("--workers", type=int, default=16)
    ap.add_argument("--redo", action="store_true")
    a = ap.parse_args()
    if a.cmd == "init":
        init(a.task)
        report(a.task)
    elif a.cmd == "pending":
        pending(a.task)
    elif a.cmd == "apply":
        apply(a.task, a.file)
        report(a.task)
    elif a.cmd == "propose":
        propose(a.task, a.workers, a.redo)
        report(a.task)
    elif a.cmd == "check":
        sys.exit(check(a.task))
    else:
        report(a.task)


if __name__ == "__main__":
    main()
