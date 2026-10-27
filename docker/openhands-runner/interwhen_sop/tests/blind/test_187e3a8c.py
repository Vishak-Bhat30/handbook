"""Blind unit tests for insurance_mojave_crest_assurance_company_187e3a8c (fixtures from the task seed only)."""

from __future__ import annotations

import json
import shutil
import subprocess

from interwhen_sop.tests.base import XLSX_PY, TaskCase

CASE = "CASE-2026-00413"
FOLDER = CASE
MEMO = f"{CASE}_Determination_Memo.docx"
PROV = "mtorres@desertspringsrecovery.com"

ACK_SUBJ = f"Mojave Crest Appeals Case #{CASE} - Acknowledgment of Your Appeal"
HEADER = ("--- EXPEDITED REVIEW ---\nExpedited Basis: Concurrent Care \u2014 Automatic (\u00a74.3.3)\n"
          "Decision Due By: 01/10/2026 11:59 PM PT\n------------------------\n\n")
OPEN = ("This letter confirms that Mojave Crest Assurance Company has received your appeal regarding continued "
        "residential substance use disorder treatment at Desert Springs Recovery Center.\n\n")
DETAILS = (f"Your case details:\nCase Number: {CASE}\nMember Name: Sarah Chen\nMember ID: VA-88291-03\n"
           "Date Received: 01/08/2026 14:23 PT\nAppeal Type: Concurrent Care\nDecision Due By: 01/10/2026 11:59 PM PT\n\n"
           "I will be your Appeals Coordinator throughout this case. If you have questions about your case, please "
           "reply to this email or call the Appeals Unit at (702) 555-0148, referencing your Case Number above.\n\n"
           f"Sincerely,\nAlex Rivera\nAppeals & Grievances Coordinator\nMojave Crest Assurance Company\n"
           f"Phone: (702) 555-0148\nCase #: {CASE}")
ACK_BODY = "Dear Dr. Torres,\n\n" + HEADER + OPEN + DETAILS

B1 = (f"CLINICAL ROUTING\nCase #: {CASE}\nType: Concurrent Care\nService Classification: SUD\n"
      "Standard of Review: Medical Necessity\nDeadline: 01/10/2026 11:59 PM PT\nParity Flag: YES\n"
      f"Case Folder: /workdir/{FOLDER}")
B2 = (f"ESCALATION\nCase #: {CASE}\nType: Concurrent\nDeadline: 01/10/2026 11:59 PM PT\nTrigger: 8, 12, 13\n"
      "Summary: Concurrent care appeal from the treating provider; coverage ends 01/10/2026 11:59 PM PT. Parity flag "
      "raised and the $85,000 dispute value exceeds Coordinator authority.\n"
      "My Recommendation: Proceed with Dr. Ontario and the Parity Review Committee; Supervisor co-signature.\n"
      "Requested Action: Approval")
B3 = (f"PARITY FLAG\nCase #: {CASE}\nService Classification: Inpatient SUD\n"
      "Comparable M/S Classification: Inpatient, in-network (inpatient rehabilitation)\n"
      "Flag Indicator: Provider asserts step therapy applied to residential SUD is not applied to M/S inpatient rehab; "
      "letter cites MHPAEA.\n"
      f"Case Folder: /workdir/{FOLDER}\nDecision Deadline: 01/10/2026")
URGENT = f"URGENT - {CASE} (concurrent care, expedited). Deadline 01/10/2026 11:59 PM PT. @dwhitfield escalation sent by DM."

LOG_ROW = {"Case ID": CASE, "Member ID": "VA-88291-03", "Group ID": "G-14827", "Plan Year": "2025",
           "Appeal Type": "Concurrent Care", "Service Classification": "Inpatient in-network (BH)",
           "Standard of Review": "Medical Necessity", "Date of Receipt": "2026-01-08 14:23",
           "Acknowledgment Due": "2026-01-09 14:23", "Acknowledgment Sent": "2026-01-08", "Tolled (Y/N)": "N",
           "Decision Due": "2026-01-10 23:59", "Reviewer Assigned": "Dr. Ontario", "Parity Flag": "Y",
           "Compliance Escalation": "N", "External Review Requested": "N", "Case Status": "With Reviewer",
           "Coordinator": "Alex Rivera"}

MEMO_SECTIONS = [
    ("Case Identification", [f"Case Number: {CASE}", "Member ID: VA-88291-03", "Plan Year: 2025", "Group ID: G-14827",
                             "Appeal Type: Concurrent Care", "Date of Receipt: 01/08/2026 14:23 PT",
                             "Decision Deadline: 01/10/2026", "Coordinator Name: PENDING", "Related Cases: None"]),
    ("Requestor and Standing", ["Requestor: Dr. Michael Torres, MD, Medical Director, Desert Springs Recovery Center "
                                "(treating provider)", "Standing Confirmed: YES", "AOR on File: N/A",
                                "Basis for Standing: PENDING"]),
    ("Disputed Service", ["Service Description: Continued residential substance use disorder treatment",
                          "CPT/HCPCS Code: H0018", "Place of Service: Residential SUD treatment facility",
                          "Billing Provider: Desert Springs Recovery Center", "Provider NPI: 1234567890",
                          "Dates of Service: 01/11/2026 to 01/15/2026", "Billed Amount: $85,000"]),
    ("Original Denial", ["Date of Original Denial: 12/12/2025", "Authorization Number: PA-2025-089234",
                         "Denial Reason Code: PENDING", "Medical Policy Cited: MP-247-v2.3-20250101",
                         "EOC Provision Cited: EOC_G-14827_2025 Section 3.2",
                         "Original Reviewer: Marisol Alvarado, RN, Utilization Management Reviewer",
                         "Denial Reasoning Summary: PENDING"]),
    ("Standard of Review Applied", ["Standard: Medical Necessity", "Rationale: PENDING"]),
    ("Clinical Review Summary", ["Reviewer Name: PENDING", "Reviewer Type: PENDING", "Date of Clinical Review: PENDING",
                                 "Clinical Finding Summary: PENDING"]),
    ("Parity Analysis", ["Parity Flag Raised at Intake: YES", "Parity Review Committee Consulted: YES",
                         "Parity Committee Concurrence Received: PENDING", "Flag Indicator Summary: PENDING",
                         "Parity Finding Summary: PENDING"]),
    ("Fraud and Mill Indicators", ["Appeal Mill Flag: NO", "Indicators Observed: 0", "Senior Coordinator Reviewed: N/A"]),
    ("Timeline Log", ["Date Received: 01/08/2026 14:23 PT", "Acknowledgment Sent: 01/08/2026", "Tolling Start: N/A",
                      "Tolling End: N/A", "Tolling Reason: N/A", "Clinical Review Received: PENDING",
                      "Parity Committee Concurrence: PENDING", "Determination Letter Sent: PENDING",
                      "Case Closed: PENDING"]),
    ("Decision and Rationale", ["Decision: PENDING", "Rationale: PENDING", "EOC Provision Applied: PENDING",
                                "Medical Policy Applied: PENDING"]),
    ("Authority", ["Total Dispute Value: $85,000", "Coordinator Authority Limit: $25,000",
                   "Within Coordinator Authority: NO", "Supervisor Approval Required: YES",
                   "Supervisor Name: Darnell Whitfield", "Supervisor Approval Obtained: PENDING"]),
    ("External Review Status", ["Case Forwarded to Canyon Ridge: NO", "Forwarding Date: N/A",
                                "Canyon Ridge Case Reference: N/A"]),
]

DOCX = """
import docx, json, sys
spec = json.loads(sys.argv[1])
d = docx.Document()
d.add_heading(spec["title"], level=1)
for sec, lines in spec["sections"]:
    if spec["bold"]:
        d.add_paragraph().add_run(sec).bold = True
    else:
        d.add_heading(sec, level=2)
    for l in lines:
        d.add_paragraph(l)
d.save(spec["path"])
"""


class Base(TaskCase):
    TASK = "insurance_mojave_crest_assurance_company_187e3a8c"

    # ---- end-state builders ----
    def log_row(self, **over):
        row = {**LOG_ROW, **over}
        self.edit("appeals_log_master.xlsx", f"row={row!r}\nws.append([row.get(h) for h in hdr])")

    def memo(self, replace=None, name=MEMO, bold=False):
        secs = [(s, [replace.get(l.split(':')[0], l) if replace else l for l in lines]) for s, lines in MEMO_SECTIONS]
        path = self.tmp / "workdir" / FOLDER / name
        path.parent.mkdir(parents=True, exist_ok=True)
        spec = {"title": f"{CASE} - Case Determination Memo", "sections": secs, "path": str(path), "bold": bold}
        subprocess.run([XLSX_PY, "-c", DOCX, json.dumps(spec)], check=True)

    def folder(self, skip=(), extra=()):
        wd = self.tmp / "workdir"
        (wd / FOLDER).mkdir(exist_ok=True)
        for f in ("denial_letter_dec12_2025.pdf", "EOC_G-14827_2025.pdf", "MP-247-v2.3-20250101.pdf",
                  "prior_auth_history.xlsx", "claims_history.xlsx") + tuple(extra):
            if f not in skip:
                shutil.copy(wd / f, wd / FOLDER / f)
        for f in ("appeal_letter_jan8_2026.pdf", "clinical_records.pdf", "parity_comparison.pdf"):
            if f not in skip:
                self.write(f"{FOLDER}/{f}", b"%PDF-1.4 seed attachment")
        if "ack" not in skip:
            self.write(f"{FOLDER}/acknowledgment_letter_A-1.txt", ACK_SUBJ + "\n\n" + ACK_BODY)

    def ack(self, body=ACK_BODY, subject=ACK_SUBJ, to=PROV):
        return "google_mail__send_email", {"to": to, "subject": subject, "body": body}

    def dm(self, uid, text):
        return "slack__send_dm", {"user_id": uid, "text": text}

    def post(self, ch, text):
        return "slack__post_message", {"channel_id": ch, "text": text}

    def send_all(self):
        for call in (self.ack(), self.dm("U002", B1), self.post("C003", URGENT), self.dm("U001", B2),
                     self.dm("U003", B3)):
            self.did(*call)

    def probs(self, rule):
        return "\n".join(p for r, p in self.done_all() if r == rule)

    def golden(self):
        self.log_row()
        self.folder()
        self.memo()
        self.send_all()


class InitialAndGolden(Base):
    def test_initial_state_lists_outstanding_obligations(self):
        got = self.done()
        self.assertEqual(set(got), {"mca.comms.messages_sent", "mca.log_row", "mca.memo", "mca.case_folder"})
        self.assertEqual(len([r for r, _ in self.done_all() if r == "mca.comms.messages_sent"]), 5)

    def test_golden_path_is_clean(self):
        self.golden()
        self.assertEqual(self.done(), {})

    def test_golden_messages_pass_pre(self):
        for call in (self.ack(), self.dm("U002", B1), self.post("C003", URGENT), self.dm("U001", B2),
                     self.dm("U003", B3)):
            self.assertEqual(self.pre(*call), set(), call[1].get("text", "")[:40])


class Ack(Base):
    def test_header_before_salutation_blocked(self):
        body = HEADER + "Dear Dr. Torres,\n\n" + OPEN + DETAILS
        self.assertIn("mca.comms.messages", self.pre(*self.ack(body)))

    def test_header_after_first_sentence_blocked(self):
        body = "Dear Dr. Torres,\n\n" + OPEN + HEADER + DETAILS
        self.assertIn("mca.comms.messages", self.pre(*self.ack(body)))

    def test_salutation_with_title_ok(self):
        body = "Dear Michael Torres, MD,\n\n" + HEADER + OPEN + DETAILS
        self.assertEqual(self.pre(*self.ack(body)), set())

    def test_wrong_deadline_blocked(self):
        body = ACK_BODY.replace("01/10/2026 11:59 PM PT", "01/11/2026 2:23 PM PT")
        self.assertIn("mca.comms.messages", self.pre(*self.ack(body)))

    def test_wrong_subject_blocked(self):
        self.assertIn("mca.comms.messages", self.pre(*self.ack(subject="Your appeal - acknowledgment")))

    def test_already_sent_blocked(self):
        self.did(*self.ack())
        self.assertIn("mca.comms.messages", self.pre(*self.ack()))

    def test_placeholder_blocked(self):
        body = ACK_BODY.replace("Alex Rivera", "[YOUR NAME]")
        self.assertIn("comms.no_placeholders", self.pre(*self.ack(body)))


class Routing(Base):
    def test_b1_in_clinical_queue_blocked(self):
        self.assertIn("mca.comms.messages", self.pre(*self.post("C002", B1)))

    def test_b1_dm_missing_parity_flag_blocked(self):
        self.assertIn("mca.comms.messages", self.pre(*self.dm("U002", B1.replace("Parity Flag: YES", "Parity Flag: NO"))))

    def test_b1_in_urgent_ok(self):
        self.assertEqual(self.pre(*self.post("C003", B1 + "\n@dwhitfield")), set())

    def test_b1_copy_to_supervisor_not_blocked(self):
        self.assertEqual(self.pre(*self.dm("U001", B1)), set())

    def test_urgent_without_tag_blocked(self):
        self.assertIn("mca.comms.messages", self.pre(*self.post("C003", f"{CASE} expedited, due 01/10/2026 23:59")))

    def test_b2_triggers(self):
        self.assertEqual(self.pre(*self.dm("U001", B2.replace("8, 12, 13", "8, 9, 12, 13"))), set())
        self.assertIn("mca.comms.messages", self.pre(*self.dm("U001", B2.replace("8, 12, 13", "12, 13"))))

    def test_b2_in_public_channel_blocked(self):
        self.assertIn("mca.phi_public", self.pre(*self.post("C001", B2)))

    def test_b3_in_parity_channel_blocked(self):
        got = self.pre(*self.post("C004", B3))
        self.assertIn("mca.parity_channel", got)

    def test_b3_wrong_classification_blocked(self):
        self.assertIn("mca.comms.messages",
                      self.pre(*self.dm("U003", B3.replace("Inpatient SUD", "Outpatient SUD"))))


class Slack(Base):
    def test_phi_in_public_channel_blocked(self):
        self.assertIn("mca.phi_public", self.pre(*self.post("C001", f"New case {CASE} for Sarah Chen logged")))
        self.assertEqual(self.pre(*self.post("C001", f"New case {CASE} logged and routed")), set())

    def test_external_mention_blocked(self):
        self.assertIn("mca.slack_external", self.pre(*self.dm("U002", B1 + "\ncc @mtorres")))


class Determination(Base):
    def test_decision_email_blocked(self):
        call = self.ack(subject=f"URGENT - Mojave Crest Appeals Case #{CASE} - Expedited Decision",
                        body="Dear Dr. Torres,\nDecision: Your appeal is APPROVED.")
        self.assertIn("mca.no_determination", self.pre(*call))

    def test_concurrence_file_needs_concurrence(self):
        call = ("syntara__writeFile", {"file_path": f"/workdir/{FOLDER}/parity_concurrence.txt", "content": "x"})
        self.did("syntara__readPDF", {"file_path": "/workdir/sop.pdf"},
                 "PARITY CONCURRENCE Case #: [CASE ID] Committee Decision: [CONCUR WITH UPHOLD]")
        self.assertIn("mca.no_determination", self.pre(*call))
        self.did("slack__get_messages", {"channel_id": "D003"},
                 f"PARITY CONCURRENCE\nCase #: {CASE}\nCommittee Decision: CONCUR WITH OVERTURN\nSafe to Issue Letter: YES")
        self.assertNotIn("mca.no_determination", self.pre(*call))

    def test_determination_letter_file_blocked(self):
        call = ("syntara__writeFile", {"file_path": f"/workdir/{FOLDER}/determination_letter_A7.docx", "content": "x"})
        self.assertIn("mca.no_determination", self.pre(*call))
        ok = ("syntara__writeFile", {"file_path": f"/workdir/{FOLDER}/acknowledgment_letter.txt", "content": "x"})
        self.assertNotIn("mca.no_determination", self.pre(*ok))


class ExternalEmail(Base):
    def test_lapse_notice_blocked(self):
        body = ACK_BODY.replace(OPEN, OPEN + "Coverage will end on 01/10/2026 unless the appeal is approved.\n")
        self.assertIn("mca.external_email", self.pre(*self.ack(body)))

    def test_medical_necessity(self):
        bad = ACK_BODY.replace(OPEN, OPEN + "Residential care is medically necessary.\n")
        ok = ACK_BODY.replace(OPEN, OPEN + "The reviewing clinician will decide whether continued care is medically "
                                           "necessary.\n")
        self.assertIn("mca.external_email", self.pre(*self.ack(bad)))
        self.assertNotIn("mca.external_email", self.pre(*self.ack(ok)))

    def test_reviewer_name_blocked(self):
        body = ACK_BODY.replace(OPEN, OPEN + "Dr. Ontario will review your case.\n")
        self.assertIn("mca.external_email", self.pre(*self.ack(body)))

    def test_subject_without_case_id_blocked(self):
        call = ("google_mail__send_email", {"to": PROV, "subject": "Status of your appeal",
                                            "body": f"Your appeal {CASE} is under review."})
        self.assertIn("mca.external_email", self.pre(*call))

    def test_regulator_blocked(self):
        call = ("google_mail__send_email", {"to": "consumer@doi.nv.gov", "subject": "Parity complaint",
                                            "body": "Notice"})
        self.assertIn("mca.no_regulator", self.pre(*call))


class Log(Base):
    def setUp(self):
        super().setUp()
        self.folder()
        self.memo()
        self.send_all()

    def test_wrong_deadline_and_reviewer(self):
        self.log_row(**{"Decision Due": "2026-01-11 14:23", "Reviewer Assigned": "RN Reviewer Team"})
        probs = [p for r, p in self.done_all() if r == "mca.log_row"]
        self.assertTrue(any("Decision Due" in p for p in probs))
        self.assertTrue(any("Reviewer Assigned" in p for p in probs))

    def test_tolled_and_pending_standing(self):
        self.log_row(**{"Tolled (Y/N)": "Y", "Case Status": "Pending Standing"})
        probs = [p for r, p in self.done_all() if r == "mca.log_row"]
        self.assertTrue(any("Tolled" in p for p in probs) and any("Case Status" in p for p in probs))

    def test_editing_existing_row_flagged(self):
        self.log_row()
        self.edit("appeals_log_master.xlsx", "ws.cell(row=3, column=col['Case Status'], value='Intake')")
        self.assertIn("mca.log_append_only", self.done())

    def test_correct_row_clean(self):
        self.log_row()
        self.assertEqual(self.done(), {})


class Memo(Base):
    def setUp(self):
        super().setUp()
        self.log_row()
        self.folder()
        self.send_all()

    def test_narrative_must_be_pending(self):
        self.memo({"Basis for Standing": "Basis for Standing: Treating provider may file concurrent care appeals."})
        self.assertIn("Basis for Standing", self.probs("mca.memo"))

    def test_wrong_values(self):
        self.memo({"Total Dispute Value": "Total Dispute Value: $85,000.00",
                   "Medical Policy Cited": "Medical Policy Cited: MP-247-v2.2-20240101"})
        probs = [p for r, p in self.done_all() if r == "mca.memo"]
        self.assertTrue(any("Total Dispute Value" in p for p in probs))
        self.assertTrue(any("Medical Policy Cited" in p for p in probs))

    def test_wrong_filename(self):
        self.memo(name=f"{CASE}_memo.docx")
        self.assertIn("mca.memo", self.done())

    def test_bold_headings_rejected(self):
        self.memo(bold=True)
        self.assertIn("Heading 2", self.probs("mca.memo"))


class Folder(Base):
    def setUp(self):
        super().setUp()
        self.log_row()
        self.send_all()

    def test_missing_documents(self):
        self.folder(skip=("claims_history.xlsx", "ack"))
        self.memo()
        p = self.probs("mca.case_folder")
        self.assertIn("claims payment history", p)
        self.assertIn("acknowledgment", p)

    def test_wrong_eoc_version(self):
        self.folder(extra=("EOC_G-14827_2026.pdf",))
        self.memo()
        self.assertIn("EOC_G-14827_2026", self.probs("mca.case_folder"))
